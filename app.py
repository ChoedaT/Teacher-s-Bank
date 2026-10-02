import streamlit as st
import tempfile
import os
import json
from io import BytesIO
from docx import Document
from pydantic import BaseModel, Field
from google import genai
from google.genai import types

import database as db

# Import curriculum mapping from curriculum.py
try:
    from curriculum import CURRICULUM_DATA
except ImportError:
    # Fallback structure if curriculum.py is missing or empty
    CURRICULUM_DATA = {}

# Initialize database
db.init_db()

st.set_page_config(page_title="Teacher's Bank", page_icon="📚", layout="wide")

# Custom Title Header
st.markdown("<h1 style='text-align: center; color: #1E3A8A; font-size: 3rem;'>Teacher's Bank</h1>", unsafe_allow_html=True)
st.markdown("<p style='text-align: center; font-size: 1.2rem; color: #4B5563;'>Central Repository & Lopon AI Revision Assistant for Educators</p>", unsafe_allow_html=True)
st.markdown("---")

# Retrieve Gemini API key from secrets or sidebar fallback
api_key = st.secrets.get("GEMINI_API_KEY") if "GEMINI_API_KEY" in st.secrets else None

if not api_key:
    api_key = st.sidebar.text_input("Enter Gemini API Key", type="password")
    if not api_key:
        st.info("👈 Please configure GEMINI_API_KEY in Streamlit Secrets or enter your key in the sidebar.")
        st.stop()

client = genai.Client(api_key=api_key)

# Pydantic Schema for Structured AI Output
class ExtractedQuestion(BaseModel):
    question_text: str = Field(description="Full text of the question, including options if multiple choice.")
    subject: str = Field(description="Subject name: English, Dzongkha, Mathematics, Science, Biology, Physics, Chemistry, History, Geography, Economics, or ICT.")
    grade_level: str = Field(description="Grade level from Grade 7 to Grade 12.")
    topic: str = Field(description="Main academic topic covered by this question.")
    subtopic: str = Field(description="Specific subtopic if applicable, or N/A.")
    marks: int = Field(description="Mark/score assigned to the question if stated, otherwise 1.")

class QuestionList(BaseModel):
    questions: list[ExtractedQuestion]

# Create Main Feature Sections
col_qb, col_lopon = st.columns(2)

with col_qb:
    st.markdown("### 📄 Question Bank")
    st.caption("Upload exam papers (PDF, MS Word, Scanned images) and save categorized questions.")

with col_lopon:
    st.markdown("### 🤖 Lopon AI")
    st.caption("Synthesize multi-school questions into structured subtopics & auto-generate revision worksheets.")

st.markdown("---")

# Main Navigation Tabs
tab_upload, tab_bank, tab_lopon = st.tabs(["📤 Upload Exam Paper", "📖 Question Bank", "🤖 Lopon AI Revision Helper"])

# ==========================================
# TAB 1: UPLOAD EXAM PAPER
# ==========================================
with tab_upload:
    st.header("Upload Assessment Paper")
    
    c1, c2, c3, c4 = st.columns(4)
    with c1:
        school_name = st.text_input("School Name", value="Phuentsholing Higher Secondary School")
    with c2:
        assessment_type = st.selectbox("Assessment Type", ["Unit Test", "Mid-Term", "Trial", "Annual", "Board Exam"])
    with c3:
        year_str = st.text_input("Year", value="2026")
    with c4:
        default_subject = st.selectbox("Subject (Hint for AI)", ["Auto-detect", "English", "Dzongkha", "Mathematics", "Science", "Biology", "Physics", "Chemistry", "History", "Geography", "Economics", "ICT"])

    uploaded_file = st.file_uploader("Upload File (PDF, MS Word .docx, PNG, JPG)", type=["pdf", "docx", "png", "jpg", "jpeg"])

    if uploaded_file and st.button("Extract & Process Questions"):
        with st.spinner("Analyzing paper with Gemini AI..."):
            try:
                suffix = os.path.splitext(uploaded_file.name)[1]
                with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp_file:
                    tmp_file.write(uploaded_file.getvalue())
                    tmp_file_path = tmp_file.name

                gemini_file = client.files.upload(file=tmp_file_path)
                curriculum_json = json.dumps(CURRICULUM_DATA, indent=2)

                prompt = f"""
                You are an expert curriculum assistant. Analyze this test document and extract every individual question.
                Categorize each question using the curriculum guidelines provided:
                - Subject (English, Dzongkha, Mathematics, Science, Biology, Physics, Chemistry, History, Geography, Economics, ICT)
                - Grade Level (Grade 7 to Grade 12)
                - Topic & Subtopic
                - Marks allocated (default to 1 if not specified)

                Subject Hint: {default_subject}
                Curriculum Topic Reference (Classes 7 to 12):
                {curriculum_json}
                """

                response = client.models.generate_content(
                    model='gemini-2.5-flash',
                    contents=[gemini_file, prompt],
                    config=types.GenerateContentConfig(
                        response_mime_type="application/json",
                        response_schema=QuestionList,
                    )
                )

                os.remove(tmp_file_path)

                parsed_data = json.loads(response.text)
                extracted_questions = parsed_data.get("questions", [])

                if extracted_questions:
                    for q in extracted_questions:
                        db.insert_question(
                            question_text=q["question_text"],
                            subject=q["subject"],
                            grade_level=q["grade_level"],
                            topic=q["topic"],
                            subtopic=q.get("subtopic", "N/A"),
                            school_name=school_name,
                            year=year_str,
                            assessment_type=assessment_type,
                            marks=q["marks"]
                        )
                    st.success(f"Successfully extracted and saved {len(extracted_questions)} questions into the bank!")
                else:
                    st.warning("No questions could be extracted. Please check the document quality.")

            except Exception as e:
                st.error(f"Error processing file: {str(e)}")

# ==========================================
# TAB 2: QUESTION BANK VIEW
# ==========================================
with tab_bank:
    st.header("Search & Filter Saved Question Papers")

    # Assessment Type Tabs
    assess_tab = st.radio("Select Assessment Category:", ["Unit Test", "Mid-Term", "Trial", "Annual", "Board Exam"], horizontal=True)

    f1, f2, f3 = st.columns(3)
    with f1:
        filter_subject = st.selectbox("Subject", ["All", "English", "Dzongkha", "Mathematics", "Science", "Biology", "Physics", "Chemistry", "History", "Geography", "Economics", "ICT"])
    with f2:
        filter_grade = st.selectbox("Grade", ["All", "Grade 7", "Grade 8", "Grade 9", "Grade 10", "Grade 11", "Grade 12"])
    with f3:
        filter_topic = st.text_input("Search Topic Keyword", value="")

    questions = db.get_questions_by_assessment(assessment_type=assess_tab, subject=filter_subject, grade=filter_grade, topic=filter_topic)

    st.subheader(f"Total Questions Found in {assess_tab}: {len(questions)}")

    for q in questions:
        q_id, q_text, q_subj, q_grade, q_topic, q_subtopic, q_school, q_year, q_marks = q
        with st.expander(f"[{q_subj} - {q_grade}] Topic: {q_topic} ({q_school}, {q_year} - {q_marks} Marks)"):
            st.write(f"**Question:** {q_text}")
            if q_subtopic and q_subtopic != "N/A":
                st.caption(f"Subtopic: {q_subtopic}")
            if st.button(f"Delete Question #{q_id}", key=f"del_{q_id}"):
                db.delete_question(q_id)
                st.rerun()

# ==========================================
# TAB 3: LOPON AI REVISION HELPER
# ==========================================
with tab_lopon:
    st.header("🤖 Lopon AI Revision Packet Synthesizer")
    st.write("Select questions from various schools to automatically synthesize and group them into subtopics for classroom revision.")

    c_s, c_g = st.columns(2)
    with c_s:
        lopon_subject = st.selectbox("Target Subject", ["English", "Dzongkha", "Mathematics", "Science", "Biology", "Physics", "Chemistry", "History", "Geography", "Economics", "ICT"], key="lopon_subj")
    with c_g:
        lopon_grade = st.selectbox("Target Grade Level", ["Grade 7", "Grade 8", "Grade 9", "Grade 10", "Grade 11", "Grade 12"], key="lopon_grd")

    available_qs = db.get_all_questions(subject=lopon_subject, grade=lopon_grade)

    if not available_qs:
        st.info("No questions found for the selected subject and grade. Please upload some papers in the Upload tab first!")
    else:
        st.subheader("Select Questions to Include in Revision Packet:")
        
        selected_q_texts = []
        for q in available_qs:
            q_id, q_text, q_subj, q_grade, q_topic, q_subtopic, q_school, q_year, q_assess, q_marks = q
            label = f"[{q_school} | {q_assess} {q_year}] ({q_topic}): {q_text[:100]}..."
            if st.checkbox(label, key=f"chk_{q_id}"):
                selected_q_texts.append(f"- Question ({q_school}, {q_assess}): {q_text} [Topic: {q_topic}]")

        if selected_q_texts and st.button("Synthesize Revision Packet with Lopon AI"):
            with st.spinner("Lopon AI is organizing questions into subtopics and creating answer guides..."):
                combined_questions = "\n".join(selected_q_texts)
                prompt_revision = f"""
                You are Lopon AI, an expert Bhutanese teaching assistant.
                Synthesize the following list of selected school exam questions into a structured classroom revision document.

                Instructions:
                1. Group the questions strictly into clear subtopics.
                2. Format each question cleanly with its original school source tag.
                3. Under each subtopic, provide a brief 'Key Revision Note' summarizing what students need to remember.
                4. Provide step-by-step model Answers & Solutions for each question at the end.

                Selected Questions:
                {combined_questions}
                """

                revision_response = client.models.generate_content(
                    model='gemini-2.5-flash',
                    contents=prompt_revision
                )

                st.markdown("### 📝 Synthesized Revision Sheet")
                st.markdown(revision_response.text)

                # Generate Word Document (.docx) Download
                doc = Document()
                doc.add_heading(f"Revision Worksheet - {lopon_subject} ({lopon_grade})", 0)
                doc.add_paragraph(f"Compiled via Lopon AI | Teacher's Bank")
                doc.add_paragraph("--------------------------------------------------")
                doc.add_paragraph(revision_response.text)

                bio = BytesIO()
                doc.save(bio)

                st.download_button(
                    label="📥 Download Revision Sheet as Word Doc (.docx)",
                    data=bio.getvalue(),
                    file_name=f"Lopon_AI_Revision_{lopon_subject}_{lopon_grade}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )

# Universal Footer Attribution
st.markdown("---")
st.markdown(
    "<p style='text-align: center; color: #6B7280; font-size: 0.9rem;'>"
    "App developed by <b>Choeda Thinley - Teacher from Phuentsholing Higher Secondary School</b>"
    "</p>", 
    unsafe_allow_html=True
)